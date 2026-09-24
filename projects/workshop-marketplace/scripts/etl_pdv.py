"""
ETL: Excel "Formulario ALTA PDV" -> prisma/pdv-seed.json (+ pdv-report.txt)

Lee la hoja ALTA (clientes activos), agrupa filas en clientes por CUIT
(detecta cadenas multi-sucursal), corrige coordenadas tomadas de "Mapa Web",
pre-clasifica canal (MAYORISTA solo si tipo == Distribuidora) y genera un
token de validacion por cliente.

Uso:  python scripts/etl_pdv.py
Salida:  prisma/pdv-seed.json  +  prisma/pdv-report.txt
"""
import json
import re
import secrets
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

XLSX = Path.home() / "Downloads" / "Formulario ALTA PDV (Respuestas) (5).xlsx"
OUT_JSON = Path(__file__).resolve().parent.parent / "prisma" / "pdv-seed.json"
OUT_REPORT = Path(__file__).resolve().parent.parent / "prisma" / "pdv-report.txt"

# Rango geografico de Argentina para validar coordenadas
LAT_MIN, LAT_MAX = -56.0, -21.0
LNG_MIN, LNG_MAX = -74.0, -53.0

PROVINCIAS_CANON = {
    "buenos aires": "Buenos Aires", "caba": "CABA",
    "ciudad autonoma de buenos aires": "CABA", "capital federal": "CABA",
    "catamarca": "Catamarca", "chaco": "Chaco", "chubut": "Chubut",
    "cordoba": "Córdoba", "corrientes": "Corrientes", "entre rios": "Entre Ríos",
    "formosa": "Formosa", "jujuy": "Jujuy", "la pampa": "La Pampa",
    "la rioja": "La Rioja", "mendoza": "Mendoza", "misiones": "Misiones",
    "neuquen": "Neuquén", "rio negro": "Río Negro", "salta": "Salta",
    "san juan": "San Juan", "san luis": "San Luis", "santa cruz": "Santa Cruz",
    "santa fe": "Santa Fe", "santiago del estero": "Santiago del Estero",
    "tierra del fuego": "Tierra del Fuego", "tucuman": "Tucumán",
}


def strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def norm_key(v) -> str:
    if v is None:
        return ""
    s = strip_accents(str(v)).lower()
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return s


def clean_str(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() in ("none", "nan", "-"):
        return None
    # Excel guarda numeros con .0 -> sacarlo cuando es un telefono/codigo
    if re.fullmatch(r"-?\d+\.0", s):
        s = s[:-2]
    return s


def digits(v) -> str:
    """Extrae solo digitos. Excel entrega numeros como float (ej. 20123456789.0);
    sin la coercion a int, str() deja un '0' de mas (el de la parte '.0')."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\D", "", str(v))


def fix_coord(v):
    """Corrige el separador de miles roto (-387.170.875 -> -38.7170875).
    Para Argentina la parte entera de lat/lng es de 2 digitos."""
    if v is None:
        return None
    s = str(v).strip().replace(" ", "")
    if s == "":
        return None
    neg = s.startswith("-")
    d = re.sub(r"\D", "", s)
    if len(d) < 3:
        return None
    num = float(d[:2] + "." + d[2:])
    return -num if neg else num


def norm_province(v):
    s = clean_str(v)
    if not s:
        return None, False
    canon = PROVINCIAS_CANON.get(norm_key(s))
    return (canon, True) if canon else (s, False)


def norm_whatsapp(v):
    """Mejor esfuerzo a formato apto para wa.me (sin +)."""
    d = digits(v)
    if not d:
        return None
    if d.startswith("549"):
        return d
    if d.startswith("54"):
        return d
    if d.startswith("0"):
        d = d[1:]
    return "549" + d if len(d) >= 8 else None


def slugify(s: str) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", strip_accents(s).lower()).strip("-")
    return base[:60] or "pdv"


def first(*vals):
    for v in vals:
        c = clean_str(v)
        if c:
            return c
    return None


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # --- indice de coordenadas desde "Mapa Web" (match por nombre) ---
    mapa = wb["Mapa Web"]
    coords = {}
    bad_coords = 0
    for r in range(2, mapa.max_row + 1):
        name = mapa.cell(r, 1).value
        if not name:
            continue
        lat = fix_coord(mapa.cell(r, 18).value)
        lng = fix_coord(mapa.cell(r, 19).value)
        if lat is None or lng is None:
            continue
        if not (LAT_MIN < lat < LAT_MAX and LNG_MIN < lng < LNG_MAX):
            bad_coords += 1
            continue
        coords[norm_key(name)] = {
            "lat": round(lat, 6), "lng": round(lng, 6),
            "redes": clean_str(mapa.cell(r, 7).value),
            "tienda": clean_str(mapa.cell(r, 8).value),
            "web": clean_str(mapa.cell(r, 9).value),
            "desc": clean_str(mapa.cell(r, 17).value),
        }

    # --- CUITs dados de baja (para reportar conflictos ALTA∩BAJA) ---
    baja = wb["BAJA"]
    baja_cuits = set()
    for r in range(2, baja.max_row + 1):
        d = digits(baja.cell(r, 8).value)
        if len(d) >= 8:
            baja_cuits.add(d)

    # --- agrupar filas de ALTA en clientes ---
    alta = wb["ALTA"]
    groups = defaultdict(list)  # key -> [row dict]
    no_name = 0
    for r in range(3, alta.max_row + 1):
        nombre_local = clean_str(alta.cell(r, 18).value)
        if not nombre_local:
            no_name += 1
            continue
        cuit = digits(alta.cell(r, 8).value)
        # clave de agrupacion: CUIT valido, si no nombre+celular
        if len(cuit) >= 8:
            key = ("cuit", cuit)
        else:
            key = ("name", norm_key(nombre_local) + "|" + digits(alta.cell(r, 11).value))
        groups[key].append({
            "nombre": nombre_local,
            "razon": clean_str(alta.cell(r, 7).value),
            "cuit": cuit,
            "tipoCuenta": clean_str(alta.cell(r, 14).value),
            "whatsapp": first(alta.cell(r, 11).value, alta.cell(r, 27).value),
            "phone": first(alta.cell(r, 26).value, alta.cell(r, 10).value),
            "email": clean_str(alta.cell(r, 13).value),
            "ig": clean_str(alta.cell(r, 31).value),
            "linkWeb": clean_str(alta.cell(r, 28).value),
            "tipoLocal": clean_str(alta.cell(r, 29).value),
            "address": first(alta.cell(r, 19).value, alta.cell(r, 9).value),
            "barrio": clean_str(alta.cell(r, 20).value),
            "cp": clean_str(alta.cell(r, 21).value),
            "localidad": clean_str(alta.cell(r, 22).value),
            "provincia": alta.cell(r, 24).value,
            "horario": clean_str(alta.cell(r, 12).value),
            "desc": clean_str(alta.cell(r, 30).value),
        })

    clientes = []
    used_slugs = set()
    rep = {
        "rows": 0, "clientes": 0, "sucursales": 0, "multi": [], "mayoristas": [],
        "sin_coords": [], "prov_no_norm": set(), "alta_en_baja": [], "wa_invalid": 0,
    }

    def uniq_slug(base):
        s = base
        i = 2
        while s in used_slugs:
            s = f"{base}-{i}"
            i += 1
        used_slugs.add(s)
        return s

    for key, rows in groups.items():
        rep["rows"] += len(rows)
        head = rows[0]
        razon = first(*[x["razon"] for x in rows]) or head["nombre"]
        whats = norm_whatsapp(first(*[x["whatsapp"] for x in rows]))
        if first(*[x["whatsapp"] for x in rows]) and not whats:
            rep["wa_invalid"] += 1
        canal = "MAYORISTA" if any((x["tipoCuenta"] or "").lower() == "distribuidora" for x in rows) else "MINORISTA"
        cuit = head["cuit"] or None

        sucursales = []
        for x in rows:
            prov, ok = norm_province(x["provincia"])
            if prov and not ok:
                rep["prov_no_norm"].add(prov)
            link = x["linkWeb"]
            is_tienda = bool(link and re.search(r"tiendanube|mercadolibre|empretienda|shop|tienda", link, re.I))
            cm = coords.get(norm_key(x["nombre"]))
            lat = cm["lat"] if cm else None
            lng = cm["lng"] if cm else None
            if lat is None:
                rep["sin_coords"].append(x["nombre"])
            ig = x["ig"]
            sucursales.append({
                "nombre": x["nombre"],
                "tipo": x["tipoLocal"] or "Otro",
                "phone": digits(x["phone"]) or None,
                "address": x["address"],
                "barrio": x["barrio"],
                "localidad": x["localidad"],
                "provincia": prov,
                "codigoPostal": x["cp"],
                "horario": x["horario"],
                "descripcion": first(x["desc"], cm["desc"] if cm else None),
                "lat": lat, "lng": lng,
                "tiendaOnline": link if is_tienda else (cm["tienda"] if cm else None),
                "web": (link if not is_tienda else None) or (cm["web"] if cm else None),
                "instagram": ig if (ig and "face" not in ig.lower()) else (cm["redes"] if cm else None),
                "facebook": ig if (ig and "face" in ig.lower()) else None,
            })

        rep["sucursales"] += len(sucursales)
        if len(sucursales) > 1:
            rep["multi"].append((razon, len(sucursales)))
        if canal == "MAYORISTA":
            rep["mayoristas"].append(razon)
        if cuit and cuit in baja_cuits:
            rep["alta_en_baja"].append(razon)

        clientes.append({
            "slug": uniq_slug(slugify(razon)),
            "validationToken": secrets.token_urlsafe(9),
            "razonSocial": razon,
            "cuit": cuit,
            "canal": canal,
            "destacado": False,
            "whatsapp": whats,
            "phone": digits(first(*[x["phone"] for x in rows])) or None,
            "email": first(*[x["email"] for x in rows]),
            "instagram": sucursales[0]["instagram"],
            "facebook": sucursales[0]["facebook"],
            "tiendaOnline": next((s["tiendaOnline"] for s in sucursales if s["tiendaOnline"]), None),
            "web": next((s["web"] for s in sucursales if s["web"]), None),
            "active": True,
            "sucursales": sucursales,
        })

    rep["clientes"] = len(clientes)
    OUT_JSON.write_text(json.dumps(clientes, ensure_ascii=False, indent=2), encoding="utf-8")

    with_coords = sum(1 for c in clientes for s in c["sucursales"] if s["lat"] is not None)
    lines = []
    lines.append("=== REPORTE ETL PDV ===")
    lines.append(f"Filas ALTA procesadas:        {rep['rows']}  (sin nombre, salteadas: {no_name})")
    lines.append(f"Clientes generados:           {rep['clientes']}")
    lines.append(f"Sucursales (pines):           {rep['sucursales']}  (con coords: {with_coords}, sin coords: {rep['sucursales']-with_coords})")
    lines.append(f"Coords descartadas (fuera AR): {bad_coords}")
    lines.append(f"WhatsApp no normalizables:    {rep['wa_invalid']}")
    lines.append("")
    lines.append(f"-- Clientes MULTI-SUCURSAL ({len(rep['multi'])}) [revisar agrupacion] --")
    for razon, n in sorted(rep["multi"], key=lambda t: -t[1]):
        lines.append(f"   {n:>3}  {razon}")
    lines.append("")
    lines.append(f"-- MAYORISTAS pre-clasificados ({len(rep['mayoristas'])}) [tipo=Distribuidora; ocultos al CF] --")
    for razon in rep["mayoristas"]:
        lines.append(f"   {razon}")
    lines.append("   NOTA: el %DTO no distingue mayoristas; revisar y marcar mas en el admin si hace falta.")
    lines.append("")
    lines.append(f"-- ALTA con CUIT tambien en BAJA ({len(rep['alta_en_baja'])}) [posible conflicto] --")
    for razon in rep["alta_en_baja"]:
        lines.append(f"   {razon}")
    lines.append("")
    lines.append(f"-- Provincias no normalizadas ({len(rep['prov_no_norm'])}) [revisar] --")
    for p in sorted(rep["prov_no_norm"]):
        lines.append(f"   {p}")
    lines.append("")
    lines.append(f"-- Sucursales SIN coordenadas ({len(rep['sin_coords'])}) [geocodificar luego] --")
    for n in rep["sin_coords"][:60]:
        lines.append(f"   {n}")
    if len(rep["sin_coords"]) > 60:
        lines.append(f"   ... y {len(rep['sin_coords'])-60} mas")
    OUT_REPORT.write_text("\n".join(lines), encoding="utf-8")

    print(f"OK -> {OUT_JSON.name}: {rep['clientes']} clientes, {rep['sucursales']} sucursales ({with_coords} con coords)")
    print(f"Reporte -> {OUT_REPORT.name}")


if __name__ == "__main__":
    main()
