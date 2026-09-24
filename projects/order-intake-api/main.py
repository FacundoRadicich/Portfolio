import os, io, base64, unicodedata, uuid, secrets, re, zipfile
from datetime import datetime, timedelta, date
from typing import Optional, Any
import pandas as pd
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Depends, Header, Request
from fastapi.responses import StreamingResponse, HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded
import anthropic

try:
    from supabase import create_client
    _SUPABASE_PKG = True
except ImportError:
    _SUPABASE_PKG = False

# ── Rate limiting ─────────────────────────────────────────────────────────────
limiter = Limiter(key_func=get_remote_address)

app = FastAPI(title="OMC Pedidos")
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Config ────────────────────────────────────────────────────────────────────
SUPABASE_URL         = os.environ.get("SUPABASE_URL", "")
SUPABASE_ANON_KEY    = os.environ.get("SUPABASE_ANON_KEY", "")
SUPABASE_SERVICE_KEY = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
ADMIN_EMAIL          = os.environ.get("ADMIN_EMAIL", "admin@example.com")
STORAGE_BUCKET       = "omc-documentos"
IVA_RATE             = 0.21  # 21% — precios del catálogo son netos (sin IVA)

LEGACY_USER = os.environ.get("BASIC_AUTH_USER", "")
LEGACY_PASS = os.environ.get("BASIC_AUTH_PASS", "")  # fallback auth only when Supabase is not configured

_sb_admin  = None
_sb_public = None

if _SUPABASE_PKG and SUPABASE_URL and SUPABASE_SERVICE_KEY:
    _sb_admin = create_client(SUPABASE_URL, SUPABASE_SERVICE_KEY)
if _SUPABASE_PKG and SUPABASE_URL and SUPABASE_ANON_KEY:
    _sb_public = create_client(SUPABASE_URL, SUPABASE_ANON_KEY)

# ── Auth ──────────────────────────────────────────────────────────────────────
_http_basic = HTTPBasic(auto_error=False)
_ALLOWED_DOC_EXTS = {"pdf", "jpg", "jpeg", "png"}
_CONTENT_TYPES = {
    "pdf": "application/pdf",
    "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "png": "image/png",
}

def _jwt_user(authorization: Optional[str]) -> Any:
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Token Bearer requerido")
    token = authorization.replace("Bearer ", "")
    try:
        return _sb_public.auth.get_user(token).user
    except Exception:
        raise HTTPException(401, "Token inválido o expirado")

def _legacy_check(credentials: Optional[HTTPBasicCredentials]):
    if not credentials:
        raise HTTPException(401, "Credenciales requeridas",
                            headers={"WWW-Authenticate": "Basic"})
    if not LEGACY_USER or not LEGACY_PASS:
        raise HTTPException(503, "Auth not configured: set SUPABASE_* or BASIC_AUTH_USER/BASIC_AUTH_PASS")
    ok = (secrets.compare_digest(credentials.username, LEGACY_USER) and
          secrets.compare_digest(credentials.password, LEGACY_PASS))
    if not ok:
        raise HTTPException(401, "Credenciales incorrectas",
                            headers={"WWW-Authenticate": "Basic"})

def require_admin(
    authorization: Optional[str] = Header(None),
    credentials: Optional[HTTPBasicCredentials] = Depends(_http_basic),
) -> dict:
    if _sb_public:
        user = _jwt_user(authorization)
        if user.email != ADMIN_EMAIL:
            raise HTTPException(403, "Solo admin")
        return {"role": "admin", "user": user, "corredor": None}
    _legacy_check(credentials)
    return {"role": "admin", "user": None, "corredor": None}

def get_caller(
    authorization: Optional[str] = Header(None),
    credentials: Optional[HTTPBasicCredentials] = Depends(_http_basic),
) -> dict:
    if _sb_public:
        user = _jwt_user(authorization)
        if user.email == ADMIN_EMAIL:
            return {"role": "admin", "user": user, "corredor": None}
        res = _sb_admin.table("corredores").select(
            "id, nombre, email, porcentaje_comision"
        ).eq("supabase_user_id", str(user.id)).execute()
        if not res.data:
            raise HTTPException(403, "No tenés acceso al sistema")
        return {"role": "corredor", "user": user, "corredor": res.data[0]}
    _legacy_check(credentials)
    return {"role": "admin", "user": None, "corredor": None}

# ── Storage helpers ───────────────────────────────────────────────────────────
def _validate_file(filename: str) -> str:
    ext = (filename or "").rsplit(".", 1)[-1].lower()
    if ext not in _ALLOWED_DOC_EXTS:
        raise HTTPException(400, f"Formato no soportado. Permitidos: {', '.join(_ALLOWED_DOC_EXTS)}")
    return ext

def _storage_upload(path: str, data: bytes, ext: str):
    ct = _CONTENT_TYPES.get(ext, "application/octet-stream")
    _sb_admin.storage.from_(STORAGE_BUCKET).upload(
        path, data, {"content-type": ct, "upsert": "true"}
    )

def _signed_url(path: str, expires: int = 3600) -> str:
    result = _sb_admin.storage.from_(STORAGE_BUCKET).create_signed_url(path, expires)
    return result.get("signedURL") or result.get("signedUrl") or result["data"]["signedUrl"]

def _assert_pedido_access(pedido_id: str, caller: dict) -> dict:
    """Retorna el pedido si el caller tiene acceso, sino 403/404."""
    res = _sb_admin.table("pedidos").select(
        "id, corredor_id, factura_url, comprobante_cobro_url, comprobante_comision_url"
    ).eq("id", pedido_id).execute()
    if not res.data:
        raise HTTPException(404, "Pedido no encontrado")
    p = res.data[0]
    if caller["role"] == "corredor" and p["corredor_id"] != caller["corredor"]["id"]:
        raise HTTPException(403, "Sin acceso a este pedido")
    return p

def _recalcular_comision(pedido_id: str):
    """Recalcula valor_neto (sin IVA) y comision_monto en base a factura y NC.
    El valor_factura se ingresa como neto (sin IVA). Comisión siempre sobre sin IVA."""
    res = _sb_admin.table("pedidos").select(
        "valor_factura, valor_nota_credito, corredor_id"
    ).eq("id", pedido_id).execute()
    if not res.data:
        return
    p = res.data[0]
    if p["valor_factura"] is None:
        return
    vf = float(p["valor_factura"])
    nc = float(p.get("valor_nota_credito") or 0)
    valor_neto = max(vf - nc, 0)                       # sin IVA
    valor_con_iva = round(valor_neto * (1 + IVA_RATE), 2)

    # Sin corredor → sin comisión
    pct = 0.0
    if p.get("corredor_id"):
        corr = _sb_admin.table("corredores").select(
            "porcentaje_comision"
        ).eq("id", p["corredor_id"]).execute()
        pct = float(corr.data[0]["porcentaje_comision"]) if corr.data else 0.0

    comision = round(valor_neto * pct / 100, 2)

    _sb_admin.table("pedidos").update({
        "valor_neto": valor_neto,
        "valor_con_iva": valor_con_iva,
        "comision_monto": comision,
    }).eq("id", pedido_id).execute()

# ── Product helpers (sin cambios) ─────────────────────────────────────────────
def sa(s):
    return ''.join(c for c in unicodedata.normalize('NFD', str(s))
                   if unicodedata.category(c) != 'Mn').lower().strip()

def get(df, r, c):
    if r >= len(df) or c >= len(df.columns):
        return None
    v = df.iloc[r, c]
    return v if pd.notna(v) else None

def qty(df, r, c):
    v = get(df, r, c)
    return int(v) if v is not None and str(v).strip() not in ['x', 'x ', '', 'nan'] else None

def find_price(price_map, nombre):
    if nombre in price_map: return price_map[nombre]
    for k, v in price_map.items():
        if sa(k) == sa(nombre): return v
    return ""

# ── Mapeos ────────────────────────────────────────────────────────────────────
TIZA_750 = {
    'marfil':    'Pintura Tiza Oh My Chalk! Marfil - 750cc',
    'stone':     'Pintura Tiza Oh My Chalk! Stone  - 750cc',
    'butter':    'Pintura Tiza Oh My Chalk! Butter  - 750cc',
    'coconut':   'Pintura Tiza Oh My Chalk! Coconut  - 750cc',
    'negro':     'Pintura Tiza Oh My Chalk! Negro  - 750cc',
    'artic blue':'Pintura Tiza Oh My Chalk! Artic Blue  - 750cc',
    'capri':     'Pintura Tiza Oh My Chalk! Capri  - 750cc',
    'grey':      'Pintura Tiza Oh My Chalk! Grey  - 750cc',
    'old pink':  'Pintura Tiza Oh My Chalk! Old Pink  - 750cc',
    'thai':      'Pintura Tiza Oh My Chalk! Thai  - 750cc',
}
COL4_SIZE_MAP = {
    ('coconut', '750'): 'Pintura Tiza Oh My Chalk! Coconut  - 750cc',
    ('coconut', '210'): 'Pintura Tiza OH My Chalk! Coconut - 210cc',
    ('coconut', '120'): 'Pintura Tiza OH My Chalk! Coconut - 120cc',
}
COL4_MAP = {
    'mandarino':'Pintura Tiza OH My Chalk! Mandarino','apricot':'Pintura Tiza OH My Chalk! Apricot',
    'berry pink':'Pintura Tiza OH My Chalk! Berry Pink','old pink':'Pintura Tiza OH My Chalk! Old Pink',
    'lila':'Pintura Tiza OH My Chalk! Lila','lila violaceo':'Pintura Tiza OH My Chalk! Lila Violaceo',
    'fucsia':'Pintura Tiza OH My Chalk! Fucsia','carmine':'Pintura Tiza OH My Chalk! Carmine',
    'cherry':'Pintura Tiza OH My Chalk! Cherry','country':'Pintura Tiza OH My Chalk! Country',
    'french blue':'Pintura Tiza OH My Chalk! French Blue',
    'artic blue':'Pintura Tiza Oh My Chalk! Artic Blue  - 750cc',
    'negro':'Pintura Tiza OH My Chalk! Negro',
    'laca jazmin':'Laca al agua color Oh My Chalk! Jazmin - 370cc',
    'laca lapacho':'Laca al agua color Oh My Chalk! Lapacho - 370cc',
    'laca jade':'Laca al agua color Oh My Chalk! Jade - 370cc',
    'laca lirio':'Laca al agua color Oh My Chalk! Lirio - 370cc',
    'laca limon':'Laca al agua color Oh My Chalk! Limon - 370cc',
}
METAL_MAP = {
    'cobre':'Pintura Tiza Metalizada OH My Chalk! Cobre - 110cc',
    'plata':'Pintura Tiza Metalizada OH My Chalk! Plata - 110cc',
    'oro':'Pintura Tiza Metalizada OH My Chalk! Oro - 110cc',
    'gold rose':'Pintura Tiza Metalizada OH My Chalk! Gold Rose - 110cc',
    'glitter':'Pintura Tiza Metalizada OH My Chalk! Glitter - 110cc',
    'champagne':'Pintura Tiza Metalizada OH My Chalk! Champagne - 110cc',
    'peltre':'Pintura Tiza Metalizada OH My Chalk! Peltre - 110cc',
    'hematite':'Pintura Tiza Metalizada OH My Chalk! Hematite - 110cc',
    'gold blue':'Pintura Tiza Metalizada OH My Chalk! Gold Blue - 110cc',
}
ACC_MAP = {
    ('satinada','grande'):'Laca al agua OH My Chalk! Satinada 370cc',
    ('satinada','mediano'):'Laca al agua OH My Chalk! Satinada 370cc',
    ('satinada','chico'):'Laca al agua OH My Chalk! Satinada 175cc',
    ('mate','grande'):'Laca al agua OH My Chalk! Mate 370cc',
    ('mate','mediano'):'Laca al agua OH My Chalk! Mate 370cc',
    ('mate','chico'):'Laca al agua OH My Chalk! Mate 175cc',
    ('extramate','grande'):'Laca al agua OH My Chalk! Extramate 370cc',
    ('extramate','mediano'):'Laca al agua OH My Chalk! Extramate 370cc',
    ('extramate','chico'):'Laca al agua OH My Chalk! Extramate 175cc',
    ('hidrolaca','grande'):'Hidrolaca al agua OH My Chalk! 1000cc',
    ('hidrolaca','mediano'):'Hidrolaca al agua OH My Chalk! 500cc',
    ('hidrolaca','chico'):'Hidrolaca al agua OH My Chalk! 175cc',
    ('bloqueador','grande'):'Bloqueador Oh My Chalk! - 750cc',
    ('bloqueador','mediano'):'Bloqueador Oh My Chalk! - 210cc',
    ('brochas','grande'):'Brocha Redonda OH My Chalk! Grande',
    ('brochas','mediano'):'Brocha Redonda OH My Chalk! Chica',
    ('cera organica','grande'):'Cera Organica Oh My Chalk!',
    ('cera organica','mediano'):'Cera Organica Oh My Chalk!',
    ('chalk paste','grande'):'Chalk Paste Oh My Chalk!  500cc',
    ('chalk paste','mediano'):'Chalk Paste Oh My Chalk!  500cc',
    ('chalk paste','chico'):'Chalk Paste Oh My Chalk! 210cc',
}
SKIP12 = {'colores','kit tela fluo','kit tela metalizado','verde fluo','naranja fluo',
          'amarillo fluo','fucsia fluo','verde metalizado','violeta metalizado',
          'azul metalizado','rosa metalizado',''}
ACC_SKIP = {'laca + accesorios','despacho','retiro local','delivery','expresso por local',
            'despacho + traslado','forma de pago','efectivo','mp qr en local','deposito',
            'transferencia','cheque / e-cheq','mercado pago (mp)',
            'envios por urbano (recomendado)','bolsas (x10)','vintage block',''}

def tiza_nombre(color, size, price_map):
    c = sa(color)
    if size == '750': return TIZA_750.get(c, f'??? Tiza 750 {c}')
    if c == 'artic blue':
        return 'Pintura Tiza OH My Chalk! Artic Blue- 210cc' if size == '210' else 'Pintura Tiza OH My Chalk! Artic Blue - 120cc'
    suffix = '- 210cc' if size == '210' else '- 120cc'
    for cand in [f'Pintura Tiza OH My Chalk! {color.strip()} {suffix}',
                 f'Pintura Tiza Oh My Chalk! {color.strip()} {suffix}']:
        for k in price_map:
            if sa(k) == sa(cand): return k
    return f'??? Tiza {size} {c}'

def col4_nombre(color_sa, size):
    base = COL4_MAP.get(color_sa)
    if not base: return f'??? col4 {color_sa}'
    if '370cc' in base or '750cc' in base: return base
    suffix = {'750':' - 750cc','210':' - 210cc','120':' - 120cc'}
    return base + suffix.get(size, ' - 210cc')

def parse_excel(df, price_map):
    acc, t750, t210, t120 = [], [], [], []
    for r in range(6, 16):
        a = get(df, r, 16)
        if not a or sa(str(a)) in ACC_SKIP: continue
        a = str(a).strip()
        for col, lbl in [(17,'grande'),(18,'mediano'),(19,'chico')]:
            q = qty(df, r, col)
            if q:
                nombre = ACC_MAP.get((sa(a), lbl), f'??? {a} {lbl}')
                acc.append((nombre, q))
    for r in range(17, 27):
        c = get(df, r, 0)
        if not c or sa(str(c)) in ['aclaraciones','']: continue
        if sa(str(c)) == 'oro antiguo': continue
        q = qty(df, r, 1)
        if q:
            nombre = METAL_MAP.get(sa(str(c)), f'??? Metal {c}')
            t210.append((nombre, q))
    for r in range(6, 16):
        c = get(df, r, 0)
        if not c or sa(str(c)) == 'metalizados': continue
        for size, col, bucket in [('750',1,t750),('210',2,t210),('120',3,t120)]:
            q = qty(df, r, col)
            if q: bucket.append((tiza_nombre(str(c).strip(), size, price_map), q))
    for r in range(6, 16):
        c = get(df, r, 4)
        if not c or sa(str(c)) in ['colores','otros','']: continue
        c_sa = sa(str(c))
        for size, col, bucket in [('750',5,t750),('210',6,t210),('120',7,t120)]:
            q = qty(df, r, col)
            if q:
                nombre = COL4_SIZE_MAP.get((c_sa, size)) or col4_nombre(c_sa, size)
                bucket.append((nombre, q))
    for r in range(6, 16):
        c = get(df, r, 8)
        if not c or sa(str(c)) in ['colores','metalizados','pintura para tela','']: continue
        for size, col, bucket in [('750',9,t750),('210',10,t210),('120',11,t120)]:
            q = qty(df, r, col)
            if q: bucket.append((tiza_nombre(str(c).strip(), size, price_map), q))
    for r in range(6, 16):
        c = get(df, r, 12)
        if not c or sa(str(c)) in SKIP12: continue
        for size, col, bucket in [('750',13,t750),('210',14,t210),('120',15,t120)]:
            q = qty(df, r, col)
            if q: bucket.append((tiza_nombre(str(c).strip(), size, price_map), q))
    return acc + t750 + t210 + t120

def hmc_to_omc(name: str, price_map: dict) -> str | None:
    """Convierte un nombre de producto en formato HMC al nombre del catálogo OMC."""
    n = name.strip()
    F = re.IGNORECASE

    m = re.search(r'HIDROLACA\s+x?(\d+)', n, F)
    if m:
        return {'1000': 'Hidrolaca al agua OH My Chalk! 1000cc',
                '500':  'Hidrolaca al agua OH My Chalk! 500cc',
                '175':  'Hidrolaca al agua OH My Chalk! 175cc',
        }.get(m.group(1), f'Hidrolaca al agua OH My Chalk! {m.group(1)}cc')

    if re.search(r'EXTRA\s*MATE', n, F):
        m = re.search(r'x?(\d+)cc', n, F); size = m.group(1) if m else '370'
        return {'370': 'Laca al agua OH My Chalk! Extramate 370cc',
                '175': 'Laca al agua OH My Chalk! Extramate 175cc'}.get(size, f'Laca al agua OH My Chalk! Extramate {size}cc')

    if re.search(r'SATINADA', n, F):
        m = re.search(r'x?(\d+)cc', n, F); size = m.group(1) if m else '370'
        return {'370': 'Laca al agua OH My Chalk! Satinada 370cc',
                '175': 'Laca al agua OH My Chalk! Satinada 175cc'}.get(size, f'Laca al agua OH My Chalk! Satinada {size}cc')

    if re.search(r'LACA\s+MATE', n, F):
        m = re.search(r'x?(\d+)cc', n, F); size = m.group(1) if m else '370'
        return {'370': 'Laca al agua OH My Chalk! Mate 370cc',
                '175': 'Laca al agua OH My Chalk! Mate 175cc'}.get(size, f'Laca al agua OH My Chalk! Mate {size}cc')

    if re.search(r'LACA CON COLOR', n, F):
        return None

    if re.search(r'BLOQUEADOR', n, F):
        m = re.search(r'x?(\d+)cc', n, F); size = m.group(1) if m else '210'
        return {'750': 'Bloqueador Oh My Chalk! - 750cc',
                '210': 'Bloqueador Oh My Chalk! - 210cc'}.get(size, f'Bloqueador Oh My Chalk! - {size}cc')

    if re.search(r'CERA ORGANICA', n, F):
        return 'Cera Organica Oh My Chalk!'

    if re.search(r'PASTE', n, F):
        m = re.search(r'x?(\d+)cc', n, F); size = m.group(1) if m else '500'
        return {'500': 'Chalk Paste Oh My Chalk!  500cc',
                '210': 'Chalk Paste Oh My Chalk! 210cc'}.get(size, f'Chalk Paste Oh My Chalk! {size}cc')

    if re.search(r'BROCHA', n, F):
        return 'Brocha Redonda OH My Chalk! Chica' if re.search(r'30mm', n, F) else 'Brocha Redonda OH My Chalk! Grande'

    if re.search(r'VINTAGE BLOCK', n, F):
        return 'Vintage Block'

    if re.search(r'PIZARRON', n, F):
        m = re.search(r'x?(\d+)cc', n, F); size = m.group(1) if m else '120'
        return f'Pintura de Tiza para Pizarron - {size}cc'

    if re.search(r'PINT\.TIZA', n, F) and re.search(r'METALGLITTER', n, F):
        return 'Pintura Tiza Metalizada OH My Chalk! Glitter - 110cc'

    if re.search(r'PINT\.TIZA', n, F) and re.search(r'METAL\s+\w', n, F):
        m = re.search(r'METAL\s+(.+)', n, F)
        if m:
            metal = m.group(1).strip().upper()
            for k, v in [('ORO ANTIGUO', None), ('GOLD ROSE', 'Pintura Tiza Metalizada OH My Chalk! Gold Rose - 110cc'),
                         ('GOLD BLUE', 'Pintura Tiza Metalizada OH My Chalk! Gold Blue - 110cc'),
                         ('COBRE', 'Pintura Tiza Metalizada OH My Chalk! Cobre - 110cc'),
                         ('PLATA', 'Pintura Tiza Metalizada OH My Chalk! Plata - 110cc'),
                         ('ORO', 'Pintura Tiza Metalizada OH My Chalk! Oro - 110cc'),
                         ('GLITTER', 'Pintura Tiza Metalizada OH My Chalk! Glitter - 110cc'),
                         ('CHAMPAGNE', 'Pintura Tiza Metalizada OH My Chalk! Champagne - 110cc'),
                         ('PELTRE', 'Pintura Tiza Metalizada OH My Chalk! Peltre - 110cc'),
                         ('HEMATITE', 'Pintura Tiza Metalizada OH My Chalk! Hematite - 110cc')]:
                if metal.startswith(k):
                    return v
        return None

    if re.search(r'PINT\.TIZA', n, F):
        m = re.search(r'x(\d+)cc\s+(.+)', n, F)
        if m:
            return tiza_nombre(m.group(2).strip().lower(), m.group(1), price_map)

    return f'??? {n}'


def parse_hmc_list(df: Any, price_map: dict) -> list:
    """Parsea un Excel de lista HMC: col0=cantidad, col1=nombre producto."""
    items = []
    for idx in range(len(df)):
        qty_val = df.iloc[idx, 0] if len(df.columns) > 0 else None
        name_val = df.iloc[idx, 1] if len(df.columns) > 1 else None
        if pd.isna(qty_val) or pd.isna(name_val):
            continue
        try:
            q = int(float(str(qty_val).strip()))
        except (ValueError, TypeError):
            continue
        if q <= 0:
            continue
        omc_name = hmc_to_omc(str(name_val).strip(), price_map)
        if omc_name:
            items.append((omc_name, q))
    return items


def build_odoo_excel(items, cliente, con_iva, price_map):
    rows = []
    for i, (nombre, q) in enumerate(items):
        precio = find_price(price_map, nombre)  # vacío si no está en catálogo; Odoo busca por nombre
        rows.append({
            "Customer*": cliente if i == 0 else "",
            "Order Date": "", "Expiration": "", "Payment Terms": "",
            "Order Lines/Products*": nombre,
            "Order Lines/Quantity": q,
            "Order Lines/Unit Price": precio,
            "Order Lines/Taxes": "IVA VENTAS" if con_iva else "",
            "Sales Team": "", "Customer Reference": "", "Tags": "",
        })
    buf = io.BytesIO()
    pd.DataFrame(rows).to_excel(buf, index=False, sheet_name='Quotations')
    buf.seek(0)
    return buf


_NOTA_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates", "nota_pedido.xlsx")


def _build_template_cell_map(price_map: dict) -> dict:
    """Returns {sa(catalog_name): (openpyxl_row, openpyxl_col)} for the OMC template."""
    m: dict = {}

    def _add(name: str, row: int, col: int):
        k = sa(name)
        if k and not k.startswith('???'):
            m[k] = (row, col)

    # Tiza Section A (cols A-D): rows 7-16
    for color, row in [
        ('Marfil', 7), ('Marfil Antiguo', 8), ('Stone', 9), ('Butter', 10),
        ('Arena', 11), ('Coconut', 12), ('Camel', 13), ('Canela', 14),
        ('África', 15), ('Vainilla', 16),
    ]:
        for size, col in [('750', 2), ('210', 3), ('120', 4)]:
            _add(tiza_nombre(color, size, price_map), row, col)

    # Tiza Section B (cols E-H): rows 7-16
    for color, row in [
        ('Mandarina', 7), ('Apricot', 8), ('Berry Pink', 9), ('Old Pink', 10),
        ('Lila', 11), ('Lila Violáceo', 12), ('Fucsia', 13), ('Carmine', 14),
        ('Cherry', 15), ('Country', 16),
    ]:
        c_sa = sa(color)
        for size, col in [('750', 6), ('210', 7), ('120', 8)]:
            _add(col4_nombre(c_sa, size), row, col)

    # Tiza Section C (cols I-L): rows 7-16
    for color, row in [
        ('Bambú', 7), ('Lemon grass', 8), ('Old grey', 9), ('Eucalipto', 10),
        ('Oliva', 11), ('Mint', 12), ('Terrenal', 13), ('Thai', 14),
        ('Bruma', 15), ('Capri', 16),
    ]:
        for size, col in [('750', 10), ('210', 11), ('120', 12)]:
            _add(tiza_nombre(color, size, price_map), row, col)

    # Tiza Section D (cols M-P): rows 7-16
    for color, row in [
        ('Sky', 7), ('Paradise', 8), ('Denim Blue', 9), ('Artic Blue', 10),
        ('French Blue', 11), ('Light Grey', 12), ('Grey', 13),
        ('French Grey', 14), ('Dark Grey', 15), ('Negro', 16),
    ]:
        for size, col in [('750', 14), ('210', 15), ('120', 16)]:
            _add(tiza_nombre(color, size, price_map), row, col)

    # Metalizados (col A, qty col B): rows 18-27
    for color, row in [
        ('Cobre', 18), ('Plata', 19), ('Oro', 20), ('Oro Antiguo', 21),
        ('Gold Rose', 22), ('Glitter', 23), ('Champagne', 24),
        ('Peltre', 25), ('Hematite', 26), ('Gold Blue', 27),
    ]:
        name = METAL_MAP.get(sa(color))
        if name:
            _add(name, row, 2)

    # Otros/Section-B overflow (cols E-H): rows 18-27
    for color, row in [
        ('French Blue', 18), ('Artic Blue', 19), ('Fucsia', 20), ('Esmeralda', 21),
        ('Negro', 22), ('Laca Jazmin', 23), ('Laca Lapacho', 24),
        ('Laca Jade', 25), ('Laca Lirio', 26), ('Laca Limon', 27),
    ]:
        c_sa = sa(color)
        base = COL4_MAP.get(c_sa)
        if not base:
            continue
        if '370cc' in base or '750cc' in base:
            _add(base, row, 6)
        else:
            for size, col in [('750', 6), ('210', 7), ('120', 8)]:
                _add(col4_nombre(c_sa, size), row, col)

    # Laca + Accesorios (cols Q-T): rows 7-16
    for name, row, col in [
        ('Laca al agua OH My Chalk! Satinada 370cc', 7, 18),
        ('Laca al agua OH My Chalk! Satinada 175cc', 7, 20),
        ('Laca al agua OH My Chalk! Mate 370cc', 8, 18),
        ('Laca al agua OH My Chalk! Mate 175cc', 8, 20),
        ('Laca al agua OH My Chalk! Extramate 370cc', 9, 18),
        ('Laca al agua OH My Chalk! Extramate 175cc', 9, 20),
        ('Hidrolaca al agua OH My Chalk! 1000cc', 10, 18),
        ('Hidrolaca al agua OH My Chalk! 500cc', 10, 19),
        ('Hidrolaca al agua OH My Chalk! 175cc', 10, 20),
        ('Bloqueador Oh My Chalk! - 750cc', 11, 18),
        ('Bloqueador Oh My Chalk! - 210cc', 11, 19),
        ('Brocha Redonda OH My Chalk! Grande', 12, 18),
        ('Brocha Redonda OH My Chalk! Chica', 12, 19),
        ('Cera Organica Oh My Chalk!', 13, 18),
        ('Chalk Paste Oh My Chalk!  500cc', 14, 18),
        ('Chalk Paste Oh My Chalk! 210cc', 14, 20),
        ('Vintage Block', 15, 18),
    ]:
        _add(name, row, col)

    return m


def build_nota_pedido(items, cliente, descuento, con_iva, price_map):
    import openpyxl

    if not os.path.exists(_NOTA_TEMPLATE):
        raise HTTPException(500, "Template de nota de pedido no encontrado en el servidor")

    wb = openpyxl.load_workbook(_NOTA_TEMPLATE)
    ws = wb['Nota de pedido']

    # Header cells (openpyxl rows are 1-based; template pandas row 1 → openpyxl row 2)
    ws.cell(row=2, column=1).value = cliente
    ws.cell(row=3, column=13).value = descuento / 100  # stored as 0.xx decimal
    ws.cell(row=4, column=1).value = date.today()

    cell_map = _build_template_cell_map(price_map)

    # Accumulate quantities (merge duplicates into same cell)
    totals: dict[tuple, int] = {}
    for name, qty_val in items:
        cell = cell_map.get(sa(name))
        if cell:
            totals[cell] = totals.get(cell, 0) + qty_val

    for (row, col), qty_val in totals.items():
        ws.cell(row=row, column=col).value = qty_val

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# ── Catálogo en memoria ───────────────────────────────────────────────────────
PRICE_MAP: dict = {}
CATALOGO_BYTES: Optional[bytes] = None

def load_catalog(data: bytes):
    global PRICE_MAP, CATALOGO_BYTES
    CATALOGO_BYTES = data
    df = pd.read_excel(io.BytesIO(data), sheet_name='Sheet1', header=0)

    cols = {str(c).strip().lower(): i for i, c in enumerate(df.columns)}
    nombre_idx = next((cols[c] for c in cols if 'nombre' in c), None)
    precio_idx = next((cols[c] for c in cols if 'precio' in c), None)

    # Formato nuevo (2 columnas: Nombre / Precio de venta) o legacy (col 4 / col 7)
    if nombre_idx is None or precio_idx is None:
        nombre_idx, precio_idx = 4, 7

    PRICE_MAP = {str(r.iloc[nombre_idx]).strip(): r.iloc[precio_idx]
                 for _, r in df.iterrows()
                 if pd.notna(r.iloc[nombre_idx]) and pd.notna(r.iloc[precio_idx])}

if os.path.exists("catalogo.xlsx"):
    try:
        with open("catalogo.xlsx", "rb") as f:
            load_catalog(f.read())
    except Exception:
        pass  # catálogo inválido o vacío — el admin lo sube desde la UI

# ── Drafts en memoria ─────────────────────────────────────────────────────────
_drafts: dict[str, dict] = {}

def _store_draft(items: list, params: dict) -> str:
    draft_id = str(uuid.uuid4())
    _drafts[draft_id] = {"items": items, "params": params, "created_at": datetime.utcnow()}
    cutoff = datetime.utcnow() - timedelta(minutes=30)
    for k in [k for k, v in _drafts.items() if v["created_at"] < cutoff]:
        del _drafts[k]
    return draft_id

def _consume_draft(draft_id: str) -> dict:
    draft = _drafts.pop(draft_id, None)
    if not draft:
        raise HTTPException(404, "Draft no encontrado o expirado")
    if datetime.utcnow() - draft["created_at"] > timedelta(minutes=30):
        raise HTTPException(404, "Draft expirado. Volvé a procesar el pedido.")
    return draft

# ── HTML helper con Cache-Control ─────────────────────────────────────────────
_NO_STORE = {"Cache-Control": "no-store, no-cache, must-revalidate", "Pragma": "no-cache"}

def _html(filename: str) -> HTMLResponse:
    path = os.path.join("static", filename)
    if not os.path.exists(path):
        raise HTTPException(404)
    with open(path, encoding="utf-8") as f:
        return HTMLResponse(f.read(), headers=_NO_STORE)

# ── Config pública ────────────────────────────────────────────────────────────
@app.get("/api/config")
def get_config():
    return {
        "supabase_url": SUPABASE_URL,
        "supabase_anon_key": SUPABASE_ANON_KEY,
        "supabase_enabled": bool(_sb_public),
    }

# ── Catálogo ──────────────────────────────────────────────────────────────────
@app.post("/api/catalogo")
async def subir_catalogo(
    file: UploadFile = File(...),
    caller: dict = Depends(require_admin),
):
    data = await file.read()
    load_catalog(data)
    with open("catalogo.xlsx", "wb") as f:
        f.write(data)
    return {"ok": True, "productos": len(PRICE_MAP)}

@app.get("/api/status")
def status(caller: dict = Depends(require_admin)):
    return {"catalogo_cargado": len(PRICE_MAP) > 0, "productos": len(PRICE_MAP)}

# ── Procesar (paso 1: draft + preview) ───────────────────────────────────────
@app.post("/api/procesar")
@limiter.limit("10/minute")
async def procesar_pedido(
    request: Request,
    file: UploadFile = File(...),
    cliente: str = Form(...),
    con_iva: bool = Form(False),
    corredor_id: str = Form(""),
    descuento: float = Form(0),
    fecha_entrega_estimada: str = Form(...),
    caller: dict = Depends(require_admin),
):
    if not PRICE_MAP:
        raise HTTPException(400, "Primero subí el catálogo de productos")
    if descuento < 0 or descuento > 100:
        raise HTTPException(400, "El descuento debe estar entre 0 y 100")

    data = await file.read()
    content_type = file.content_type or ""
    filename = file.filename or ""
    items = []

    if content_type.startswith("image/") or filename.lower().endswith((".jpg",".jpeg",".png",".webp")):
        ext = filename.split(".")[-1].lower()
        media_types = {"jpg":"image/jpeg","jpeg":"image/jpeg","png":"image/png","webp":"image/webp"}
        media_type = media_types.get(ext, "image/jpeg")
        b64 = base64.standard_b64encode(data).decode()
        api_key = os.environ.get("ANTHROPIC_API_KEY")
        if not api_key:
            raise HTTPException(503, "El procesamiento de fotos no está disponible (falta ANTHROPIC_API_KEY).")
        try:
            client = anthropic.Anthropic(api_key=api_key)
            resp = client.messages.create(
                model="claude-sonnet-4-20250514",
                max_tokens=1000,
                messages=[{
                    "role": "user",
                    "content": [
                        {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
                        {"type": "text", "text": (
                            "Esta es una foto de un pedido manuscrito de pinturas Oh My Chalk. "
                            "Extraé los productos y cantidades. "
                            "Respondé SOLO con líneas en formato: CANTIDAD|PRODUCTO "
                            "Por ejemplo: 6|Marfil 210cc\n3|Butter 750cc\n "
                            "Interpretá los colores aunque estén mal escritos. "
                            "Si dice el tamaño (750, 210, 120) incluilos, si no, asumí 210cc."
                        )}
                    ]
                }]
            )
            text = resp.content[0].text.strip()
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(502, f"Error al procesar la imagen con IA: {e}")
        for line in text.split("\n"):
            line = line.strip()
            if "|" not in line: continue
            parts = line.split("|", 1)
            try:
                q = int(parts[0].strip())
                desc = parts[1].strip()
                size = "210"
                if "750" in desc: size = "750"
                elif "120" in desc: size = "120"
                color = desc.replace("750cc","").replace("210cc","").replace("120cc","").strip()
                items.append((tiza_nombre(color, size, PRICE_MAP), q))
            except Exception:
                continue

    elif filename.lower().endswith((".xlsx",".xls")):
        try:
            xl = pd.ExcelFile(io.BytesIO(data))
        except Exception:
            raise HTTPException(400, "No se pudo abrir el Excel. Verificá que el archivo no esté dañado.")
        try:
            if 'Nota de pedido' in xl.sheet_names:
                df = pd.read_excel(xl, sheet_name='Nota de pedido', header=None)
                items = parse_excel(df, PRICE_MAP)
            else:
                # Formato lista (ej: HMC): col0=cantidad, col1=nombre
                first_sheet = xl.sheet_names[0]
                df = pd.read_excel(xl, sheet_name=first_sheet, header=None)
                if len(df.columns) < 2:
                    raise HTTPException(
                        400,
                        f"Formato no reconocido. Hojas: {', '.join(xl.sheet_names)}. "
                        "Usá la plantilla estándar 'Nota de pedido' o una imagen del pedido."
                    )
                items = parse_hmc_list(df, PRICE_MAP)
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(400, f"Error al leer la nota de pedido: {e}")
    else:
        raise HTTPException(400, "Formato no soportado. Subí un Excel (.xlsx) o una imagen (.jpg, .png)")

    if not items:
        raise HTTPException(400, "No se encontraron productos en el archivo")

    # Calcular subtotal automáticamente (precio × cantidad de productos con precio conocido)
    items_preview = []
    subtotal = 0.0
    sin_precio = 0
    for n, q in items:
        precio = find_price(PRICE_MAP, n)
        precio_num = float(precio) if isinstance(precio, (int, float)) else 0.0
        if precio_num > 0:
            subtotal += precio_num * q
        else:
            sin_precio += 1
        items_preview.append({"nombre": n, "cantidad": q, "precio": precio})

    valor_sin_iva = round(subtotal * (1 - descuento / 100), 2)   # base de comisión
    valor_con_iva = round(valor_sin_iva * (1 + IVA_RATE), 2)

    draft_id = _store_draft(items, {
        "cliente": cliente, "con_iva": con_iva,
        "corredor_id": corredor_id or None,
        "descuento": descuento,
        "subtotal": round(subtotal, 2),
        "valor_estimado": valor_sin_iva,
        "fecha_entrega_estimada": fecha_entrega_estimada,
    })
    return {
        "draft_id": draft_id, "cliente": cliente, "total_items": len(items),
        "items": items_preview, "subtotal": round(subtotal, 2),
        "descuento": descuento,
        "valor_sin_iva": valor_sin_iva,
        "valor_con_iva": valor_con_iva,
        "iva_rate": IVA_RATE,
        "sin_precio": sin_precio,
    }

# ── Confirmar (paso 2: crea en Supabase + devuelve Excel) ─────────────────────
@app.post("/api/confirmar")
def confirmar_pedido(
    draft_id: str = Form(...),
    valor_estimado: Optional[float] = Form(None),  # override editable (valor SIN IVA) desde la preview
    caller: dict = Depends(require_admin),
):
    draft = _consume_draft(draft_id)
    items = draft["items"]
    p = draft["params"]

    cliente   = p["cliente"]
    con_iva   = p.get("con_iva", False)
    descuento = p.get("descuento", 0)

    buf_odoo = build_odoo_excel(items, cliente, con_iva, PRICE_MAP)
    # The order-form template is company-specific and not shipped in this repo
    buf_nota = (build_nota_pedido(items, cliente, descuento, con_iva, PRICE_MAP)
                if os.path.exists(_NOTA_TEMPLATE) else None)

    # Valor SIN IVA (base de comisión): el editado en la preview, o el calculado
    valor_sin_iva = valor_estimado if valor_estimado is not None and valor_estimado > 0 else p["valor_estimado"]
    valor_con_iva = round(valor_sin_iva * (1 + IVA_RATE), 2)

    if _sb_admin:
        corredor_id = p.get("corredor_id")
        comision_inicial = 0.0
        if corredor_id:
            corr_res = _sb_admin.table("corredores").select(
                "porcentaje_comision"
            ).eq("id", corredor_id).execute()
            pct = float(corr_res.data[0]["porcentaje_comision"]) if corr_res.data else 7.0
            comision_inicial = round(valor_sin_iva * pct / 100, 2)
        _sb_admin.table("pedidos").insert({
            "corredor_id": corredor_id,
            "cliente": cliente,
            "valor_estimado": valor_sin_iva,
            "valor_con_iva": valor_con_iva,
            "con_iva": con_iva,
            "descuento": descuento,
            "fecha_entrega_estimada": p["fecha_entrega_estimada"],
            "estado": "en_preparacion",
            "comision_monto": comision_inicial,
        }).execute()

    # ZIP con los dos archivos
    nombre_base = cliente.replace(' ', '_')[:30]
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{nombre_base}_odoo.xlsx", buf_odoo.read())
        if buf_nota:
            zf.writestr(f"{nombre_base}_nota_pedido.xlsx", buf_nota.read())
    zip_buf.seek(0)

    return StreamingResponse(
        zip_buf,
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename=pedido_{nombre_base}.zip"}
    )

# ── Documentos: Factura ───────────────────────────────────────────────────────
@app.post("/api/pedidos/{pedido_id}/factura")
async def subir_factura(
    pedido_id: str,
    file: UploadFile = File(...),
    valor_factura: float = Form(...),
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    if valor_factura <= 0:
        raise HTTPException(400, "El valor de la factura debe ser mayor a 0")
    ext = _validate_file(file.filename or "")
    data = await file.read()
    path = f"facturas/{pedido_id}/factura.{ext}"
    _storage_upload(path, data, ext)
    _sb_admin.table("pedidos").update({
        "factura_url": path,
        "valor_factura": valor_factura,
    }).eq("id", pedido_id).execute()
    _recalcular_comision(pedido_id)
    return {"ok": True}

@app.get("/api/pedidos/{pedido_id}/factura/url")
def get_factura_url(pedido_id: str, caller: dict = Depends(get_caller)):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    p = _assert_pedido_access(pedido_id, caller)
    if not p.get("factura_url"):
        raise HTTPException(404, "No hay factura cargada")
    return {"url": _signed_url(p["factura_url"])}

# ── Documentos: Nota de Crédito ───────────────────────────────────────────────
@app.post("/api/pedidos/{pedido_id}/nota-credito")
async def subir_nota_credito(
    pedido_id: str,
    file: UploadFile = File(...),
    valor_nota_credito: float = Form(...),
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    if valor_nota_credito < 0:
        raise HTTPException(400, "El valor de la nota de crédito no puede ser negativo")
    ext = _validate_file(file.filename or "")
    data = await file.read()
    path = f"notas-credito/{pedido_id}/nc.{ext}"
    _storage_upload(path, data, ext)
    _sb_admin.table("pedidos").update({
        "nota_credito_url": path,
        "valor_nota_credito": valor_nota_credito,
    }).eq("id", pedido_id).execute()
    _recalcular_comision(pedido_id)
    return {"ok": True}

@app.get("/api/pedidos/{pedido_id}/nota-credito/url")
def get_nc_url(pedido_id: str, caller: dict = Depends(get_caller)):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    res = _sb_admin.table("pedidos").select(
        "id, corredor_id, nota_credito_url"
    ).eq("id", pedido_id).execute()
    if not res.data: raise HTTPException(404, "Pedido no encontrado")
    p = res.data[0]
    if caller["role"] == "corredor" and p["corredor_id"] != caller["corredor"]["id"]:
        raise HTTPException(403, "Sin acceso")
    if not p.get("nota_credito_url"):
        raise HTTPException(404, "No hay nota de crédito cargada")
    return {"url": _signed_url(p["nota_credito_url"])}

# ── Documentos: Comprobante de Cobro (corredor) ───────────────────────────────
@app.post("/api/pedidos/{pedido_id}/comprobante-cobro")
async def subir_comprobante_cobro(
    pedido_id: str,
    file: UploadFile = File(...),
    fecha_cobro: str = Form(...),
    forma_pago: str = Form(...),
    caller: dict = Depends(get_caller),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    formas_validas = ("echeq", "cheque", "transferencia")
    if forma_pago not in formas_validas:
        raise HTTPException(400, f"Para subir comprobante la forma de pago debe ser: {formas_validas}")

    _assert_pedido_access(pedido_id, caller)
    ext = _validate_file(file.filename or "")
    data = await file.read()
    path = f"comprobantes-cobro/{pedido_id}/comprobante.{ext}"
    _storage_upload(path, data, ext)

    _sb_admin.table("pedidos").update({
        "comprobante_cobro_url": path,
        "fecha_cobro": fecha_cobro,
        "forma_pago": forma_pago,
        "cobro_pendiente_validacion": True,
    }).eq("id", pedido_id).execute()
    return {"ok": True}

@app.get("/api/pedidos/{pedido_id}/comprobante-cobro/url")
def get_comprobante_cobro_url(pedido_id: str, caller: dict = Depends(get_caller)):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    p = _assert_pedido_access(pedido_id, caller)
    if not p.get("comprobante_cobro_url"):
        raise HTTPException(404, "No hay comprobante de cobro cargado")
    return {"url": _signed_url(p["comprobante_cobro_url"])}

# ── Cobro en efectivo (admin valida directo) ──────────────────────────────────
@app.post("/api/pedidos/{pedido_id}/cobro-efectivo")
def registrar_cobro_efectivo(
    pedido_id: str,
    fecha_cobro: str = Form(...),
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    res = _sb_admin.table("pedidos").select(
        "fecha_entrega_real"
    ).eq("id", pedido_id).execute()
    dias_al_cobro = None
    if res.data and res.data[0]["fecha_entrega_real"]:
        entrega = date.fromisoformat(res.data[0]["fecha_entrega_real"])
        cobro = date.fromisoformat(fecha_cobro)
        dias_al_cobro = (cobro - entrega).days

    _sb_admin.table("pedidos").update({
        "fecha_cobro": fecha_cobro,
        "forma_pago": "efectivo",
        "estado": "cobrado",
        "cobro_pendiente_validacion": False,
        "dias_al_cobro": dias_al_cobro,
    }).eq("id", pedido_id).execute()
    return {"ok": True, "dias_al_cobro": dias_al_cobro}

# ── Validar cobro con comprobante (admin) ─────────────────────────────────────
@app.patch("/api/pedidos/{pedido_id}/validar-cobro")
def validar_cobro(
    pedido_id: str,
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    res = _sb_admin.table("pedidos").select(
        "fecha_cobro, fecha_entrega_real"
    ).eq("id", pedido_id).execute()
    if not res.data:
        raise HTTPException(404, "Pedido no encontrado")
    p = res.data[0]
    dias_al_cobro = None
    if p["fecha_entrega_real"] and p["fecha_cobro"]:
        entrega = date.fromisoformat(p["fecha_entrega_real"])
        cobro = date.fromisoformat(p["fecha_cobro"])
        dias_al_cobro = (cobro - entrega).days

    _sb_admin.table("pedidos").update({
        "estado": "cobrado",
        "cobro_pendiente_validacion": False,
        "dias_al_cobro": dias_al_cobro,
    }).eq("id", pedido_id).execute()
    return {"ok": True, "dias_al_cobro": dias_al_cobro}

# ── Documentos: Comprobante de Comisión (admin) ───────────────────────────────
@app.post("/api/pedidos/{pedido_id}/comprobante-comision")
async def subir_comprobante_comision(
    pedido_id: str,
    file: UploadFile = File(...),
    comision_fecha_pago: str = Form(...),
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    ext = _validate_file(file.filename or "")
    data = await file.read()
    path = f"comprobantes-comision/{pedido_id}/comprobante.{ext}"
    _storage_upload(path, data, ext)
    _sb_admin.table("pedidos").update({
        "comprobante_comision_url": path,
        "comision_estado": "pagada",
        "comision_fecha_pago": comision_fecha_pago,
    }).eq("id", pedido_id).execute()
    return {"ok": True}

@app.get("/api/pedidos/{pedido_id}/comprobante-comision/url")
def get_comprobante_comision_url(pedido_id: str, caller: dict = Depends(get_caller)):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    p = _assert_pedido_access(pedido_id, caller)
    if not p.get("comprobante_comision_url"):
        raise HTTPException(404, "No hay comprobante de comisión cargado")
    return {"url": _signed_url(p["comprobante_comision_url"])}

# ── Comisión pagada en efectivo (sin comprobante) ─────────────────────────────
@app.patch("/api/pedidos/{pedido_id}/comision-efectivo")
def comision_efectivo(
    pedido_id: str,
    comision_fecha_pago: str = Form(...),
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    _sb_admin.table("pedidos").update({
        "comision_estado": "pagada",
        "comision_fecha_pago": comision_fecha_pago,
    }).eq("id", pedido_id).execute()
    return {"ok": True}

# ── Pedidos CRUD ──────────────────────────────────────────────────────────────
@app.get("/api/pedidos")
def listar_pedidos(caller: dict = Depends(get_caller)):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    q = _sb_admin.table("pedidos").select(
        "id, corredor_id, cliente, valor_estimado, valor_factura, valor_nota_credito, valor_neto, "
        "fecha_pedido, fecha_entrega_estimada, fecha_entrega_real, estado, "
        "comprobante_cobro_url, cobro_pendiente_validacion, fecha_cobro, forma_pago, dias_al_cobro, "
        "comision_monto, comision_estado, comision_fecha_pago, "
        "factura_url, nota_credito_url, comprobante_comision_url, "
        "corredores(nombre, email)"
    ).order("fecha_pedido", desc=True)

    if caller["role"] == "corredor":
        q = q.eq("corredor_id", caller["corredor"]["id"])

    return q.execute().data

@app.get("/api/pedidos/{pedido_id}")
def get_pedido(pedido_id: str, caller: dict = Depends(get_caller)):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    res = _sb_admin.table("pedidos").select(
        "id, corredor_id, cliente, valor_estimado, valor_factura, valor_nota_credito, valor_neto, "
        "fecha_pedido, fecha_entrega_estimada, fecha_entrega_real, estado, "
        "comprobante_cobro_url, cobro_pendiente_validacion, fecha_cobro, forma_pago, dias_al_cobro, "
        "comision_monto, comision_estado, comision_fecha_pago, notas, "
        "factura_url, nota_credito_url, comprobante_comision_url, "
        "corredores(nombre, email)"
    ).eq("id", pedido_id).execute()
    if not res.data:
        raise HTTPException(404, "Pedido no encontrado")
    p = res.data[0]
    if caller["role"] == "corredor" and p["corredor_id"] != caller["corredor"]["id"]:
        raise HTTPException(403, "Sin acceso a este pedido")
    return p

@app.patch("/api/pedidos/{pedido_id}/estado")
def actualizar_estado(
    pedido_id: str,
    estado: str = Form(...),
    fecha_entrega_real: Optional[str] = Form(None),
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    if estado not in ("en_preparacion", "entregado", "cobrado"):
        raise HTTPException(400, "Estado inválido")
    update: dict = {"estado": estado}
    if estado == "entregado" and fecha_entrega_real:
        update["fecha_entrega_real"] = fecha_entrega_real
    _sb_admin.table("pedidos").update(update).eq("id", pedido_id).execute()
    return {"ok": True}

@app.patch("/api/pedidos/{pedido_id}/comision")
def actualizar_comision(
    pedido_id: str,
    comision_estado: str = Form(...),
    comision_fecha_pago: Optional[str] = Form(None),
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    if comision_estado not in ("pendiente", "aprobada", "pagada"):
        raise HTTPException(400, "Estado de comisión inválido")
    update: dict = {"comision_estado": comision_estado}
    if comision_estado == "pagada" and comision_fecha_pago:
        update["comision_fecha_pago"] = comision_fecha_pago
    _sb_admin.table("pedidos").update(update).eq("id", pedido_id).execute()
    return {"ok": True}

# ── Corredores ────────────────────────────────────────────────────────────────
@app.get("/api/corredores")
def listar_corredores(caller: dict = Depends(require_admin)):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    res = _sb_admin.table("corredores").select(
        "id, nombre, email, porcentaje_comision"
    ).eq("activo", True).order("nombre").execute()
    return res.data

@app.post("/api/corredores")
def crear_corredor(
    nombre: str = Form(...),
    email: str = Form(...),
    porcentaje_comision: float = Form(7.0),
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    nombre = nombre.strip()
    email  = email.strip().lower()
    if not nombre:
        raise HTTPException(400, "El nombre es obligatorio")
    existing = _sb_admin.table("corredores").select("id").eq("email", email).execute()
    if existing.data:
        raise HTTPException(409, "Ya existe un corredor con ese email")
    res = _sb_admin.table("corredores").insert({
        "nombre": nombre,
        "email":  email,
        "porcentaje_comision": porcentaje_comision,
        "activo": True,
    }).execute()
    return res.data[0]

@app.delete("/api/corredores/{corredor_id}")
def desactivar_corredor(
    corredor_id: str,
    caller: dict = Depends(require_admin),
):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    _sb_admin.table("corredores").update({"activo": False}).eq("id", corredor_id).execute()
    return {"ok": True}

# ── Dashboard resumen ─────────────────────────────────────────────────────────
@app.get("/api/dashboard/resumen")
def dashboard_resumen(caller: dict = Depends(get_caller)):
    if not _sb_admin:
        raise HTTPException(503, "Supabase no configurado")
    q = _sb_admin.table("pedidos").select(
        "estado, comision_monto, comision_estado, fecha_entrega_real, fecha_cobro, cobro_pendiente_validacion"
    )
    if caller["role"] == "corredor":
        q = q.eq("corredor_id", caller["corredor"]["id"])
    pedidos = q.execute().data

    hoy = date.today()
    comision_pendiente = comision_pagada = 0.0
    en_riesgo = vencidos = cobros_a_validar = 0

    for p in pedidos:
        cm = float(p["comision_monto"] or 0)
        if p["comision_estado"] == "pagada":
            comision_pagada += cm
        else:
            comision_pendiente += cm

        if p.get("cobro_pendiente_validacion"):
            cobros_a_validar += 1

        if p["estado"] == "entregado" and not p["fecha_cobro"] and p["fecha_entrega_real"]:
            dias = (hoy - date.fromisoformat(p["fecha_entrega_real"])).days
            if dias >= 30:   vencidos += 1
            elif dias >= 20: en_riesgo += 1

    return {
        "comision_pendiente": round(comision_pendiente, 2),
        "comision_pagada": round(comision_pagada, 2),
        "en_riesgo": en_riesgo,
        "vencidos": vencidos,
        "cobros_a_validar": cobros_a_validar,
        "total_pedidos": len(pedidos),
    }

# ── Páginas HTML ──────────────────────────────────────────────────────────────
@app.get("/", response_class=HTMLResponse)
def index(): return _html("index.html")

@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(): return _html("dashboard.html")

@app.get("/admin", response_class=HTMLResponse)
def admin_panel(): return _html("admin.html")

app.mount("/static", StaticFiles(directory="static"), name="static")
